"""
patch_ppt_slides.py — surgically rebuild ONLY slides 10, 22, 23 of an existing
deck (keeps every other slide, incl. the user's hand-edited slide 8):
  - Slide 10: pipeline as a flow chart (boxes + arrows)
  - Slide 22 (Next Steps): three filled blocks
  - Slide 23 (Takeaways): 2x2 grid of blocks
Matches the deck's navy+teal theme.

    .venv/bin/python scripts/patch_ppt_slides.py "<path to .pptx>"
"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

INK   = RGBColor(0x12, 0x22, 0x39); BAND = RGBColor(0x1B, 0x30, 0x4C)
ACCENT= RGBColor(0x17, 0xB3, 0xA3); AMBER= RGBColor(0xF0, 0x8A, 0x24)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
BG    = RGBColor(0xF5, 0xF7, 0xFA); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SLATE = RGBColor(0x30, 0x3E, 0x4E); MUTE = RGBColor(0x8A, 0x97, 0xA6)
LINEC = RGBColor(0xDD, 0xE3, 0xE9); LIGHT = RGBColor(0xE7, 0xF6, 0xF3)

path = sys.argv[1]
prs = Presentation(path)
SW, SH = prs.slide_width, prs.slide_height
MARGIN, BAND_H, FOOT_Y = Inches(0.55), Inches(1.12), Inches(7.05)


def kill_shadow(shape):
    spPr = shape._element.spPr
    for el in spPr.findall(qn("a:effectLst")):
        spPr.remove(el)
    spPr.append(spPr.makeelement(qn("a:effectLst"), {}))
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def rect(slide, l, t, w, h, fill, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, l, t, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.fill.background(); s.shadow.inherit = False; kill_shadow(s)
    return s


def clear(slide):
    for sh in list(slide.shapes):
        sh._element.getparent().remove(sh._element)
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = BG


def title_band(slide, title):
    rect(slide, 0, 0, SW, BAND_H, BAND)
    rect(slide, 0, BAND_H, SW, Pt(6), ACCENT)
    tb = slide.shapes.add_textbox(MARGIN, 0, SW - 2 * MARGIN, BAND_H)
    tf = tb.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    r = tf.paragraphs[0].add_run(); r.text = title
    r.font.size = Pt(27); r.font.bold = True; r.font.color.rgb = WHITE


def footer(slide, n):
    tb = slide.shapes.add_textbox(SW - MARGIN - Inches(1.2), FOOT_Y, Inches(1.2), Inches(0.32))
    p = tb.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
    r = p.add_run(); r.text = str(n); r.font.size = Pt(11); r.font.color.rgb = MUTE


def lead(slide, text, y=Inches(1.55)):
    tb = slide.shapes.add_textbox(MARGIN, y, SW - 2 * MARGIN, Inches(0.6))
    r = tb.text_frame.paragraphs[0].add_run(); r.text = text
    r.font.size = Pt(18); r.font.color.rgb = SLATE


def block(slide, l, t, w, h, header, body, color, hsize=16, bsize=13.5):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    box.fill.solid(); box.fill.fore_color.rgb = color
    box.line.fill.background(); box.shadow.inherit = False; kill_shadow(box)
    tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.2); tf.margin_right = Inches(0.2)
    p1 = tf.paragraphs[0]; p1.alignment = PP_ALIGN.CENTER
    r1 = p1.add_run(); r1.text = header; r1.font.bold = True; r1.font.size = Pt(hsize); r1.font.color.rgb = WHITE
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER; p2.space_before = Pt(6)
    r2 = p2.add_run(); r2.text = body; r2.font.size = Pt(bsize); r2.font.color.rgb = LIGHT


def right_arrow(slide, l, t, w, h, color):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, l, t, w, h)
    a.fill.solid(); a.fill.fore_color.rgb = color
    a.line.fill.background(); a.shadow.inherit = False; kill_shadow(a)


# ── Slide 10 — Pipeline flow chart ────────────────────────────────────────────
s = prs.slides[9]
clear(s); title_band(s, "Approach: the Pipeline"); footer(s, 10)
lead(s, "A scripted expert teaches; at run-time the neural network is fully in control.")
stages = [("Scripted Expert", "hand-coded IK controller", ACCENT),
          ("Demonstrations", "240 episodes · 3 cameras + language", BAND),
          ("Fine-tune SmolVLA", "freeze backbone, train action head", ACCENT),
          ("Closed-loop Eval", "success rate, not loss", BAND)]
bw, gap, bh, top = Inches(2.55), Inches(0.62), Inches(1.95), Inches(3.0)
x0 = int((SW - (4 * bw + 3 * gap)) / 2)
xs = [x0 + i * (bw + gap) for i in range(4)]
for i, (hd, sub, col) in enumerate(stages):
    block(s, xs[i], top, bw, bh, hd, sub, col, hsize=15, bsize=12)
ay = top + int(bh / 2) - Inches(0.28)
for i in range(3):
    right_arrow(s, xs[i] + bw + Inches(0.06), ay, gap - Inches(0.12), Inches(0.56), RGBColor(0x45, 0x5A, 0x6A))
cap = s.shapes.add_textbox(MARGIN, Inches(5.5), SW - 2 * MARGIN, Inches(0.7))
cp = cap.text_frame.paragraphs[0]; cp.alignment = PP_ALIGN.CENTER
cr = cp.add_run(); cr.text = "The expert generates clean demos → the policy learns to imitate → then it drives on its own."
cr.font.size = Pt(15); cr.font.italic = True; cr.font.color.rgb = MUTE

# ── Slide 22 — Next Steps as three blocks ─────────────────────────────────────
s = prs.slides[21]
clear(s); title_band(s, "Next Steps"); footer(s, 22)
lead(s, "This is Part A — the manipulation half. What comes next:")
nb = [("NAVIGATION VLA", "Drive the base between shelves in a store with moving, colliding shoppers.", ACCENT),
      ("ORCHESTRATOR", "Loop the full list:  navigate → pick → place → deliver.", BAND),
      ("ROBUSTNESS", "A small retrain with non-empty baskets → reliable multi-item lists.", AMBER)]
bw, gap, bh, top = Inches(3.75), Inches(0.55), Inches(3.0), Inches(2.75)
x0 = int((SW - (3 * bw + 2 * gap)) / 2)
for i, (hd, bd, col) in enumerate(nb):
    block(s, x0 + i * (bw + gap), top, bw, bh, hd, bd, col)

# ── Slide 23 — Takeaways as a 2x2 grid of blocks ──────────────────────────────
s = prs.slides[22]
clear(s); title_band(s, "Takeaways"); footer(s, 23)
tk = [("Runs on 12 GB", "A capable VLA fine-tuned and deployed on a single consumer GPU.", BAND),
      ("80% task completion", "Rigorously measured — 320 trials, Wilson CIs, success-based selection.", GREEN),
      ("Understands language", "Reaches the named item; a live window into the model's 50-step plan.", ACCENT),
      ("End-to-end system", "A typed shopping list → items collected in the basket.", AMBER)]
bw, bh, gx, gy = Inches(5.75), Inches(2.0), Inches(0.5), Inches(0.4)
x0 = int((SW - (2 * bw + gx)) / 2); y0 = Inches(1.75)
for i, (hd, bd, col) in enumerate(tk):
    r, c = divmod(i, 2)
    block(s, x0 + c * (bw + gx), y0 + r * (bh + gy), bw, bh, hd, bd, col, hsize=18, bsize=14)

prs.save(path)
print(f"patched slides 10, 22, 23 in {path}")
