"""
patch_add_slides.py — add two slides to an existing deck without disturbing the rest:
  - "If There Were No Constraints — the Ideal System"  (inserted after Next Steps)
  - "Thank You"  (appended last)
Then renumbers all footer page-numbers so they stay consistent.

    .venv/bin/python scripts/patch_add_slides.py "<path to .pptx>"
"""
import re
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

INK   = RGBColor(0x12, 0x22, 0x39); BAND = RGBColor(0x1B, 0x30, 0x4C)
ACCENT= RGBColor(0x17, 0xB3, 0xA3); AMBER= RGBColor(0xF0, 0x8A, 0x24)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
BG    = RGBColor(0xF5, 0xF7, 0xFA); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SLATE = RGBColor(0x30, 0x3E, 0x4E); MUTE = RGBColor(0x8A, 0x97, 0xA6)
LIGHT = RGBColor(0xE7, 0xF6, 0xF3)

path = sys.argv[1]
prs = Presentation(path)
SW, SH = prs.slide_width, prs.slide_height
MARGIN, BAND_H, FOOT_Y = Inches(0.55), Inches(1.12), Inches(7.05)
BLANK = next((l for l in prs.slide_layouts if l.name == "Blank"), prs.slide_layouts[6])


def kill_shadow(shape):
    spPr = shape._element.spPr
    for el in spPr.findall(qn("a:effectLst")):
        spPr.remove(el)
    spPr.append(spPr.makeelement(qn("a:effectLst"), {}))
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def rect(slide, l, t, w, h, fill, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, l, t, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.fill.background(); s.shadow.inherit = False; kill_shadow(s)
    return s


def slide_bg(slide, color=BG):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = color


def title_band(slide, title):
    rect(slide, 0, 0, SW, BAND_H, BAND)
    rect(slide, 0, BAND_H, SW, Pt(6), ACCENT)
    tb = slide.shapes.add_textbox(MARGIN, 0, SW - 2 * MARGIN, BAND_H)
    tf = tb.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = True
    r = tf.paragraphs[0].add_run(); r.text = title
    r.font.size = Pt(26); r.font.bold = True; r.font.color.rgb = WHITE


def block(slide, l, t, w, h, header, body, color, hsize=16, bsize=13.5):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    box.fill.solid(); box.fill.fore_color.rgb = color
    box.line.fill.background(); box.shadow.inherit = False; kill_shadow(box)
    tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.2); tf.margin_right = Inches(0.2)
    p1 = tf.paragraphs[0]; p1.alignment = PP_ALIGN.CENTER
    r1 = p1.add_run(); r1.text = header; r1.font.bold = True; r1.font.size = Pt(hsize); r1.font.color.rgb = WHITE
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER; p2.space_before = Pt(6)
    r2 = p2.add_run(); r2.text = body; r2.font.size = Pt(bsize); r2.font.color.rgb = LIGHT


# ── build the two new slides (appended; reordered below) ──────────────────────
# Ideal / unconstrained system
s_ideal = prs.slides.add_slide(BLANK)
slide_bg(s_ideal); title_band(s_ideal, "If There Were No Constraints — the Ideal System")
lead = s_ideal.shapes.add_textbox(MARGIN, Inches(1.45), SW - 2 * MARGIN, Inches(0.55))
lr = lead.text_frame.paragraphs[0].add_run()
lr.text = "With unlimited compute, the answer isn't a bigger model — it's the right architecture:"
lr.font.size = Pt(17); lr.font.color.rgb = SLATE
# model callout bar
cbar = rect(s_ideal, MARGIN, Inches(2.05), SW - 2 * MARGIN, Inches(0.62), BAND, MSO_SHAPE.ROUNDED_RECTANGLE)
ctf = cbar.text_frame; ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
cp = ctf.paragraphs[0]; cp.alignment = PP_ALIGN.CENTER
c0 = cp.add_run(); c0.text = "Model:  "; c0.font.size = Pt(15); c0.font.bold = True; c0.font.color.rgb = ACCENT
c1 = cp.add_run(); c1.text = "a dual-system generalist (π0.5 / GR00T-class) — a large reasoner + a fast action policy"
c1.font.size = Pt(15); c1.font.color.rgb = WHITE
# 2x2 design blocks
blocks = [("UNIFIED WHOLE-BODY", "Drive base + arm together — reach while repositioning, not park-then-pick.", ACCENT),
          ("TACTILE + DEPTH SENSING", "Touch fixes grasping (the milk / cereal weakness); depth + LiDAR for the drop & the crowd.", AMBER),
          ("DIVERSE CO-TRAINED DATA", "Teleop across stores + randomized sim + human video → generalizes to any store, not 4 items.", GREEN),
          ("IMITATION → REINFORCEMENT", "RL fine-tuning pushes past the demonstration ceiling.", BAND)]
bw, bh, gx, gy = Inches(5.75), Inches(1.5), Inches(0.5), Inches(0.3)
x0 = int((SW - (2 * bw + gx)) / 2); y0 = Inches(3.0)
for i, (hd, bd, col) in enumerate(blocks):
    r, c = divmod(i, 2)
    block(s_ideal, x0 + c * (bw + gx), y0 + r * (bh + gy), bw, bh, hd, bd, col, hsize=15, bsize=12.5)
botb = s_ideal.shapes.add_textbox(MARGIN, Inches(6.35), SW - 2 * MARGIN, Inches(0.6))
bp = botb.text_frame.paragraphs[0]; bp.alignment = PP_ALIGN.CENTER
br0 = bp.add_run(); br0.text = "My SmolVLA system is the deployable subset of this ideal"
br0.font.size = Pt(15); br0.font.bold = True; br0.font.color.rgb = INK
br1 = bp.add_run(); br1.text = " — same recipe, shrunk to one GPU."
br1.font.size = Pt(15); br1.font.color.rgb = SLATE

# Thank you
s_thanks = prs.slides.add_slide(BLANK)
rect(s_thanks, 0, 0, SW, SH, INK)
rect(s_thanks, int(SW / 2 - Inches(1.4)), Inches(4.25), Inches(2.8), Pt(4), ACCENT)
tt = s_thanks.shapes.add_textbox(Inches(1), Inches(2.4), SW - Inches(2), Inches(1.8))
ttf = tt.text_frame; ttf.vertical_anchor = MSO_ANCHOR.MIDDLE
tp = ttf.paragraphs[0]; tp.alignment = PP_ALIGN.CENTER
tr = tp.add_run(); tr.text = "Thank You"; tr.font.size = Pt(58); tr.font.bold = True; tr.font.color.rgb = WHITE
qb = s_thanks.shapes.add_textbox(Inches(1), Inches(4.5), SW - Inches(2), Inches(0.6))
qp = qb.text_frame.paragraphs[0]; qp.alignment = PP_ALIGN.CENTER
qr = qp.add_run(); qr.text = "Questions & discussion welcome"; qr.font.size = Pt(22); qr.font.color.rgb = ACCENT
cb = s_thanks.shapes.add_textbox(Inches(1), Inches(5.9), SW - Inches(2), Inches(0.8))
cptf = cb.text_frame
c1p = cptf.paragraphs[0]; c1p.alignment = PP_ALIGN.CENTER
c1r = c1p.add_run(); c1r.text = "Biswajeet Panda"; c1r.font.size = Pt(18); c1r.font.bold = True; c1r.font.color.rgb = WHITE
c2p = cptf.add_paragraph(); c2p.alignment = PP_ALIGN.CENTER; c2p.space_before = Pt(5)
c2r = c2p.add_run(); c2r.text = "Supervisors:  Dr Gilbert Tang   ·   Prof Phil Webb"
c2r.font.size = Pt(14); c2r.font.color.rgb = MUTE

# ── reorder: put the Ideal slide right after "Next Steps" (index 21) ───────────
sldIdLst = prs.slides._sldIdLst
kids = list(sldIdLst)
ideal_el = kids[-2]           # s_ideal (second-to-last after both appends)
takeaways_el = kids[22]       # original Takeaways
sldIdLst.remove(ideal_el)
takeaways_el.addprevious(ideal_el)

# ── renumber footers: position+1, skip title (first) and Thank You (last) ─────
slides = list(prs.slides)
last = slides[-1]
for i, s in enumerate(slides):
    if i == 0 or s is last:
        continue
    # remove any existing bottom-right bare-number textbox
    for sh in list(s.shapes):
        if sh.has_text_frame and re.fullmatch(r"\d{1,2}", sh.text_frame.text.strip() or "") \
           and sh.left is not None and sh.left > Inches(10.5) and sh.top is not None and sh.top > Inches(6.6):
            sh._element.getparent().remove(sh._element)
    tb = s.shapes.add_textbox(SW - MARGIN - Inches(1.2), FOOT_Y, Inches(1.2), Inches(0.32))
    p = tb.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
    r = p.add_run(); r.text = str(i + 1); r.font.size = Pt(11); r.font.color.rgb = MUTE

prs.save(path)
print(f"added 2 slides + renumbered; total now {len(prs.slides._sldIdLst)} slides -> {path}")
